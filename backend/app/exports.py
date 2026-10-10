"""Exportadores: Excel (tablas VLSM + evaluación) y Word (documento explicativo)."""
from io import BytesIO

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .advisor import narrative
from .schemas import Project

V4_HEAD = ["Piso/Área", "VLAN", "Usuarios", "Requeridos (con crecimiento)", "Red", "Prefijo", "Máscara",
           "Gateway", "Primera IP", "Última IP", "Broadcast", "Hosts útiles"]


def build_xlsx(p: Project, rows: list[dict], cls: dict, ev: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "VLSM IPv4"
    ws.append(V4_HEAD)
    for r in rows:
        ws.append([r["name"], r["vlan"], r["users"], r["required"], r["network"], f"/{r['prefix']}", r["mask"],
                   r["gateway"], r["first"], r["last"], r["broadcast"], r["usable"]])
    w6 = wb.create_sheet("IPv6")
    w6.append(["Piso/Área", "VLAN", "Subred /64", "Gateway"])
    for r in rows:
        w6.append([r["name"], r["vlan"], r["v6"], r["v6_gw"]])
    rs = wb.create_sheet("Evaluación")
    rs.append(["Proyecto", p.name]); rs.append(["Tipo de red", cls["type"]])
    rs.append(["Éxito del proyecto", f"{ev['score']}%"]); rs.append(["Veredicto", ev["verdict"]]); rs.append([])
    rs.append(["Severidad", "Área", "Problema", "Mejora", "Penalización"])
    for f in ev["findings"]:
        rs.append([f["severity"], f["area"], f["issue"], f["fix"], f["penalty"]])
    for sheet in wb.worksheets:                      # estilo de encabezados y anchos
        for c in sheet[1 if sheet.title != "Evaluación" else 6]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0A3D38")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()


def build_docx(p: Project, rows: list[dict], cls: dict, ev: dict) -> bytes:
    d = Document()
    d.add_heading(f"Infraestructura de red — {p.name}", 0)
    d.add_heading("Resumen", 1)
    for para in narrative(p, cls, ev):
        d.add_paragraph(para)
    d.add_heading("Direccionamiento", 1)
    t = d.add_table(rows=1, cols=6); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Piso/Área", "VLAN", "Red IPv4", "Gateway", "Hosts", "IPv6 /64"]):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate([r["name"], r["vlan"], f"{r['network']}/{r['prefix']}", r["gateway"], r["usable"], r["v6"]]):
            cells[i].text = str(v)
    d.add_heading("Aspectos a mejorar", 1)
    for f in ev["findings"]:
        d.add_paragraph(f"[{f['severity'].upper()}] {f['area']}: {f['issue']} Mejora: {f['fix']}", style="List Bullet")
    buf = BytesIO(); d.save(buf)
    return buf.getvalue()
