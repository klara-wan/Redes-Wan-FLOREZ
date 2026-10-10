"""Pruebas reales de NetPlan. Ejecutar desde /backend:  pytest -v
Cada prueba tiene un identificador [P#] que se reporta en docs/pruebas.md."""
from io import BytesIO
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app import advisor, configs, engine, exports
from app.schemas import Project

FLOORS = [{"name": n, "users": u} for n, u in [("Piso 1", 200), ("Piso 2", 150), ("Piso 3", 130), ("Piso 4", 120),
          ("Piso 5", 100), ("Piso 6", 90), ("Piso 7", 80), ("Piso 8", 60), ("Piso 9", 40), ("Piso 10", 30)]]


def proj(**kw):
    base = dict(name="Sede A", base_cidr="172.16.0.0/21", floors=FLOORS, vendor="cisco", growth_pct=0)
    base.update(kw)
    return Project(**base)


def run(p):
    rows = engine.plan_vlsm(p); cls = engine.classify(p, rows)
    return rows, cls, advisor.evaluate(p, rows, cls)


def test_p01_subnetting_sede_a():
    rows = engine.plan_vlsm(proj())
    assert (rows[0]["network"], rows[0]["prefix"]) == ("172.16.0.0", 24)
    assert (rows[3]["network"], rows[3]["prefix"], rows[3]["broadcast"]) == ("172.16.3.0", 25, "172.16.3.127")
    assert (rows[9]["network"], rows[9]["prefix"], rows[9]["usable"]) == ("172.16.5.128", 27, 30)

def test_p02_total_direcciones_1440():
    assert sum(r["usable"] + 2 for r in engine.plan_vlsm(proj())) == 1440

def test_p03_rechaza_cidr_invalido():
    with pytest.raises(ValidationError): proj(base_cidr="999.1.1.1/24")

def test_p04_rechaza_direccion_que_no_es_red():
    with pytest.raises(ValidationError): proj(base_cidr="172.16.0.5/21")

def test_p05_rechaza_ipv6_como_base():
    with pytest.raises(ValidationError): proj(base_cidr="2001:db8::/32")

def test_p06_bloque_insuficiente():
    with pytest.raises(ValueError): engine.plan_vlsm(proj(base_cidr="172.16.0.0/24"))

def test_p07_ipv6_64_por_vlan():
    rows = engine.plan_vlsm(proj())
    assert rows[0]["v6"] == "2001:db8:1:1::/64" and rows[9]["v6"] == "2001:db8:1:a::/64"

def test_p08_rechaza_prefijo_ipv6_malicioso():
    with pytest.raises(ValidationError): proj(v6_prefix="2001:db8;rm -rf")

def test_p09_clasifica_pequena():
    p = proj(floors=[{"name": "Oficina", "users": 20}], base_cidr="192.168.1.0/24")
    assert run(p)[1]["type"] == "pequeña"

def test_p10_clasifica_empresarial():
    assert run(proj())[1]["type"] == "empresarial"

def test_p11_redundancia_mejora_puntaje():
    assert run(proj(redundancy=True, dual_isp=True, growth_pct=20))[2]["score"] > run(proj())[2]["score"]

def test_p12_puntaje_nunca_100_con_secretos_de_ejemplo():
    assert run(proj(redundancy=True, dual_isp=True, growth_pct=20))[2]["score"] < 100

@pytest.mark.parametrize("vendor", ["cisco", "huawei", "fortinet", "mikrotik", "juniper", "aruba"])
def test_p13_config_por_marca_incluye_vlan_ip_y_host(vendor):
    p = proj(vendor=vendor, growth_pct=0)
    rows, cls, _ = run(p)
    devs = configs.build_devices(p, rows, cls)
    text = "\n".join(devs.values())
    assert "172.16.0.1" in text and "vlan" in text.lower() and ("CORE" in text)

def test_p14_config_cisco_exacta():
    p = proj(); rows, cls, _ = run(p)
    core = configs.build_devices(p, rows, cls)["CORE"]
    assert "hostname CORE" in core and "vlan 10" in core and "ip address 172.16.0.1 255.255.255.0" in core

def test_p15_sin_secretos_reales_en_configs():
    for v in ["cisco", "huawei", "fortinet", "mikrotik"]:
        p = proj(vendor=v); rows, cls, _ = run(p)
        text = "\n".join(configs.build_devices(p, rows, cls).values())
        assert "CAMBIAR_" in text and "admin123" not in text and "password cisco" not in text

def test_p16_no_habilita_telnet_ni_http_cisco():
    p = proj(); rows, cls, _ = run(p)
    t = configs.build_devices(p, rows, cls)["CORE"]
    assert "transport input ssh" in t and "no ip http server" in t and "transport input telnet" not in t

def test_p17_excel_tiene_hojas_vlsm():
    p = proj(); rows, cls, ev = run(p)
    wb = load_workbook(BytesIO(exports.build_xlsx(p, rows, cls, ev)))
    assert wb.sheetnames == ["VLSM IPv4", "IPv6", "Evaluación"]

def test_p18_docx_valido():
    p = proj(); rows, cls, ev = run(p)
    assert exports.build_docx(p, rows, cls, ev)[:2] == b"PK"
